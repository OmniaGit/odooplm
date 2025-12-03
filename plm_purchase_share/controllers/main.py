# -*- coding: utf-8 -*-
import io
import re
import zipfile
import requests
import base64
from odoo import http
from odoo.http import request
from openpyxl import load_workbook


class PortalPurchaseDownload(http.Controller):

    @http.route('/my/purchase/<int:order_id>/download_docs', type='http', auth='public', website=True)
    def download_purchase_documents(self, order_id, **kwargs):
        order = request.env['purchase.order'].sudo().browse(order_id)
        if not order or not order.exists():
            return request.not_found()

        zip_buffer = io.BytesIO()
        zip_file = zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED)

        report_action = request.env.ref(
            'plm.report_product_product_pdf_latest'
        ).sudo()

        Report = request.env['ir.actions.report'].sudo()

        for line in order.order_line:
            product = line.product_id
            if not product:
                continue

            html_content, _ = report_action._render_qweb_html("plm.report_product_product_pdf_latest", product.id)

            wrapped_html = f"""
                            <html>
                                <main>
                                    {html_content}
                                </main>
                            </html>
                        """

            bodies = Report._prepare_html(wrapped_html)[0]

            pdf_content = Report._run_wkhtmltopdf(bodies)

            filename = f"{product.display_name.replace('/', '_')}.pdf"
            zip_file.writestr(filename, pdf_content)

        zip_file.close()
        zip_buffer.seek(0)

        zip_filename = f"PO_{order.name}_Documents.zip"
        return request.make_response(
            zip_buffer.getvalue(),
            headers=[
                ('Content-Type', 'application/zip'),
                ('Content-Disposition', f'attachment; filename=\"{zip_filename}\"'),
            ]
        )


def clean_filename(name):
    """
    Remove or replace characters that cause zip subfolders.
    Allowed: A-Z a-z 0-9 _ - .
    Everything else → _
    """
    name = name.replace(" ", "_")
    return re.sub(r"[^A-Za-z0-9_\-\.]", "_", name)


class PurchaseZipDownload(http.Controller):

    @http.route(
        '/my/purchase/<int:order_id>/download_docs/<int:product_id>',
        type='http',
        auth="user",
        website=True
    )
    def download_zip(self, order_id, product_id, **kwargs):

        order = request.env['purchase.order'].sudo().browse(order_id)
        product = request.env['product.product'].sudo().browse(product_id)
        tmpl = product.product_tmpl_id

        if tmpl.eng_product_type != "BG":
            return request.redirect('/my/purchase/%s' % order.id)

        xls_action = tmpl.with_context(
            active_model='product.product',
            active_id=product.id
        ).act_get_xlsx_bom_report()

        download_url = xls_action.get("url")
        if not download_url:
            return request.not_found()

        base_url = request.env["ir.config_parameter"].sudo().get_param("web.base.url")
        full_url = f"{base_url}{download_url}" if download_url.startswith("/") else download_url

        # Maintain session for authenticated download
        session = requests.Session()
        for k, v in request.httprequest.cookies.items():
            session.cookies.set(k, v)

        resp = session.get(full_url, stream=True)
        if resp.status_code != 200:
            return request.not_found()

        in_file = io.BytesIO(resp.content)
        wb = load_workbook(filename=in_file)

        for sheet_name in ["Purchase Parts list", "Bolts list"]:
            if sheet_name in wb.sheetnames:
                wb.remove(wb[sheet_name])

        out_xls_buffer = io.BytesIO()
        wb.save(out_xls_buffer)
        out_xls_buffer.seek(0)
        final_xls_bytes = out_xls_buffer.read()

        xls_filename = f"{clean_filename(tmpl.name)}_bom.xlsx"

        report_action = request.env.ref('plm.report_product_product_pdf_latest').sudo()

        html_content, _ = report_action._render_qweb_html(
            "plm.report_product_product_pdf_latest",
            product.id
        )

        wrapped_html = f"<html><main>{html_content}</main></html>"
        Report = request.env['ir.actions.report'].sudo()
        body = Report._prepare_html(wrapped_html)[0]
        main_pdf = Report._run_wkhtmltopdf(body)

        main_pdf_filename = f"{clean_filename(tmpl.name)}_drawing.pdf"

        # ---------------------------------------------------------
        # 4) ALL BOM COMPONENT PDFs
        # ---------------------------------------------------------
        bom = request.env['mrp.bom'].sudo().search([
            ('product_tmpl_id', '=', tmpl.id)
        ], limit=1)

        bom_pdfs = []

        if bom:
            exploded_lines = bom.explode(product, 1)[1]

            for line in exploded_lines:
                comp = line[0]

                if not comp or not comp.exists():
                    continue

                comp_html, _ = report_action._render_qweb_html(
                    "plm.report_product_product_pdf_latest",
                    comp.id
                )

                wrapped = f"<html><main>{comp_html}</main></html>"
                comp_body = Report._prepare_html(wrapped)[0]
                comp_pdf = Report._run_wkhtmltopdf(comp_body)

                comp_name = clean_filename(comp.product_tmpl_id.name)
                filename = f"{comp_name}_drawing.pdf"

                bom_pdfs.append((filename, comp_pdf))

        zip_buffer = io.BytesIO()

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zipf:

            zipf.writestr(xls_filename, final_xls_bytes)

            zipf.writestr(f"BG/{main_pdf_filename}", main_pdf)

            for filename, pdf_content in bom_pdfs:
                zipf.writestr(f"BG/{filename}", pdf_content)

        zip_buffer.seek(0)

        # ---------------------------------------------------------
        # 6) ZIP FILE NAME
        # ---------------------------------------------------------
        vendor = clean_filename(order.partner_id.name) if order.partner_id else "_"
        project_no = clean_filename(order.project_id.name) if order.project_id else "_"
        purchase_no = clean_filename(order.name) if order.name else "_"

        zip_name = f"{project_no}-{purchase_no}.{vendor}.zip"

        # ---------------------------------------------------------
        # 7) RETURN ZIP RESPONSE
        # ---------------------------------------------------------
        return request.make_response(
            zip_buffer.read(),
            headers=[
                ("Content-Type", "application/zip"),
                ("Content-Disposition", f'attachment; filename={zip_name}')
            ]
        )
